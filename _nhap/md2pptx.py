#!/usr/bin/env python3
"""Chuyển các file .md bài giảng 28Tech thành slide PowerPoint (.pptx).

Dùng:
  python3 md2pptx.py md-cpp ppt-cpp "C++ Cơ Bản"
  python3 md2pptx.py md-stl ppt-stl "STL C++" --merge
  python3 md2pptx.py md-cpp ppt-cpp "C++ Cơ Bản" --only 04-01-cpp-vong-lap-for.md

Bố cục dòng chảy: nội dung đổ liên tiếp từ trên xuống, hết chỗ mới sang slide mới.
Đoạn dẫn ngắn (vd "Cú pháp :") luôn đi cùng code/ảnh ngay sau nó.
"""
import re, sys, glob, math
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE as _S
from pptx.util import Length
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE

W, H = Inches(13.333), Inches(7.5)
NAVY   = RGBColor(0x14, 0x2B, 0x45)
BLUE   = RGBColor(0x35, 0x98, 0xDB)
INK    = RGBColor(0x1F, 0x2A, 0x37)
GREY   = RGBColor(0x6B, 0x72, 0x80)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
CODEBG = RGBColor(0x1E, 0x29, 0x3B)
CODEFG = RGBColor(0xE6, 0xED, 0xF3)
FONT, MONO = "Calibri", "Consolas"

MARGIN      = Inches(0.7)
TOP_IN      = 1.42
BOTTOM_IN   = 7.5 - 0.45
BODY_W      = W - MARGIN * 2
BODY_H_IN   = BOTTOM_IN - TOP_IN

PARA_PT, LIST_PT = 16.5, 15.5
PARA_CPL, LIST_CPL = 88, 85
LINE_IN = 0.275
CODE_MAX_LINES = 21

def IN(v): return Emu(int(v * 914400))

# ================= parse =================
def parse(path):
    raw = open(path, encoding="utf-8").read()
    title = re.search(r"^# (.+)$", raw, re.M).group(1).strip()
    m = re.search(r"^> Nguồn: (.+)$", raw, re.M)
    src = m.group(1).strip() if m else ""
    body = re.sub(r"^# .+$\n?", "", raw, count=1, flags=re.M)
    body = re.sub(r"^> Nguồn:.+$\n?", "", body, count=1, flags=re.M)

    blocks, lines, i = [], body.split("\n"), 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            buf = []; i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i]); i += 1
            blocks.append(["code", "\n".join(buf).rstrip()]); i += 1
        elif ln.startswith("!["):
            mm = re.match(r"!\[\]\((.+)\)", ln)
            if mm: blocks.append(["img", mm.group(1)])
            i += 1
        elif ln.strip() == "---":
            blocks.append(["hr", ""]); i += 1
        elif re.match(r"^\s*(?:[-*]|\d+\.)\s+", ln):
            items = []
            while i < len(lines) and re.match(r"^\s*(?:[-*]|\d+\.)\s+", lines[i]):
                items.append(re.sub(r"^\s*(?:[-*]|\d+\.)\s+", "", lines[i]).strip()); i += 1
            blocks.append(["list", items])
        elif ln.strip():
            buf = []
            while i < len(lines) and lines[i].strip() and not lines[i].startswith(("```", "![", "---")) \
                    and not re.match(r"^\s*(?:[-*]|\d+\.)\s+", lines[i]):
                buf.append(lines[i].strip()); i += 1
            blocks.append(["para", " ".join(buf)])
        else:
            i += 1

    sections, cur = [], {"title": None, "blocks": []}
    for kind, val in blocks:
        if kind == "hr":
            if cur["blocks"] or cur["title"]:
                sections.append(cur); cur = {"title": None, "blocks": []}
            continue
        mm = re.match(r"^\*\*\s*(\d+)\s*[\.\)]?\s*(.+?)\s*\*\*$", val.strip()) if kind == "para" else None
        if mm:
            if cur["blocks"] or cur["title"]: sections.append(cur)
            cur = {"title": f"{mm.group(1)}. {mm.group(2)}", "blocks": []}
        else:
            cur["blocks"].append([kind, val])
    if cur["blocks"] or cur["title"]: sections.append(cur)
    return title, src, sections

def group_units(blocks):
    """Ghép đoạn dẫn ngắn với code/ảnh ngay sau thành 1 đơn vị không tách rời."""
    units, i = [], 0
    while i < len(blocks):
        kind, val = blocks[i]
        if kind == "para" and i + 1 < len(blocks) and blocks[i+1][0] in ("code", "img") and len(val) <= 170:
            units.append([blocks[i], blocks[i+1]]); i += 2
        else:
            units.append([blocks[i]]); i += 1
    return units

# ================= đo =================
def code_font(n, maxlen=0):
    fs = 14.0 if n <= 12 else (12.5 if n <= 17 else 11.5)
    if maxlen:                       # 1384 / fs = số ký tự vừa bề ngang
        while fs > 8.5 and maxlen > 1384 / fs:
            fs -= 0.5
    return fs

def code_metrics(text):
    lines = text.split("\n")
    return len(lines), max((len(l) for l in lines), default=0)

def img_ratio(f):
    try:
        from PIL import Image
        iw, ih = Image.open(f).size
        return ih / iw
    except Exception:
        return 0.6

def height_of(block, root):
    kind, val = block
    if kind == "para":
        n = max(1, math.ceil(len(re.sub(r"[*`]", "", val)) / PARA_CPL))
        return n * LINE_IN + 0.14
    if kind == "list":
        h = 0.0
        for it in val:
            h += max(1, math.ceil(len(re.sub(r"[*`]", "", it)) / LIST_CPL)) * (LINE_IN - 0.01) + 0.08
        return h + 0.08
    if kind == "code":
        n, mx = code_metrics(val)
        return n * (code_font(n, mx) * 1.42 / 72) + 0.34
    if kind == "img":
        f = (root / val).resolve()
        if not f.exists(): return 0.0
        w = BODY_W / 914400
        h = w * img_ratio(f)
        return min(h, BODY_H_IN - 0.1) + 0.12
    return 0.0

# ================= vẽ =================
def rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh

def rich(p, text, size, color, bold=False):
    text = re.sub(r"\[(.*?)\]\([^)]*\)", r"\1", text)
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not part: continue
        r = p.add_run()
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r.text = part[2:-2]; r.font.bold = True; r.font.name = FONT; r.font.color.rgb = color
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r.text = part[1:-1]; r.font.name = MONO; r.font.color.rgb = BLUE
        else:
            r.text = part; r.font.name = FONT; r.font.color.rgb = color
        r.font.size = Pt(size)
        if bold: r.font.bold = True

class Deck:
    def __init__(self, prs, kicker):
        self.prs, self.kicker = prs, kicker
        self.slide, self.y, self.title, self.cont = None, 0.0, "", False

    def start_section(self, title):
        self.title, self.cont, self.slide = title, False, None

    def _new(self):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        rect(s, 0, 0, W, H, WHITE)
        t = self.title if not self.cont else f"{self.title} (tiếp)"
        rect(s, MARGIN, Inches(0.58), Inches(0.085), Inches(0.5), BLUE)
        tb = s.shapes.add_textbox(MARGIN + Inches(0.24), Inches(0.44), W - MARGIN*2 - Inches(0.3), Inches(0.8))
        tf = tb.text_frame; tf.word_wrap = True
        rich(tf.paragraphs[0], t, 25, NAVY, bold=True)
        kb = s.shapes.add_textbox(MARGIN, Inches(0.11), W - MARGIN*2, Inches(0.28))
        kp = kb.text_frame.paragraphs[0]; kp.alignment = PP_ALIGN.RIGHT
        r = kp.add_run(); r.text = self.kicker
        r.font.size = Pt(10.5); r.font.color.rgb = GREY; r.font.name = FONT
        self.slide, self.y, self.cont = s, TOP_IN, True
        return s

    def ensure(self, need):
        if self.slide is None or self.y + need > BOTTOM_IN:
            self._new()

    def para(self, text, h):
        self.ensure(h)
        tb = self.slide.shapes.add_textbox(MARGIN, IN(self.y), BODY_W, IN(h))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        rich(p, text, PARA_PT, INK)
        self.y += h

    def bullets(self, items, h):
        self.ensure(h)
        tb = self.slide.shapes.add_textbox(MARGIN, IN(self.y), BODY_W, IN(h))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT; p.space_after = Pt(5)
            r = p.add_run(); r.text = "•  "
            r.font.size = Pt(LIST_PT); r.font.color.rgb = BLUE; r.font.name = FONT
            rich(p, it, LIST_PT, INK)
        self.y += h

    def code(self, text, h):
        lines = text.split("\n"); fs = code_font(*code_metrics(text))
        self.ensure(h)
        rect(self.slide, MARGIN, IN(self.y), BODY_W, IN(h - 0.1), CODEBG)
        tb = self.slide.shapes.add_textbox(MARGIN + Inches(0.2), IN(self.y + 0.12),
                                           BODY_W - Inches(0.4), IN(h - 0.32))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT; p.space_after = Pt(0); p.line_spacing = 1.0
            r = p.add_run(); r.text = ln if ln.strip() else " "
            r.font.name = MONO; r.font.size = Pt(fs); r.font.color.rgb = CODEFG
        self.y += h

    def image(self, path, h, root):
        f = (root / path).resolve()
        if not f.exists(): return
        self.ensure(h)
        ratio = img_ratio(f)
        w_in = BODY_W / 914400
        if w_in * ratio > h - 0.12:
            w_in = (h - 0.12) / ratio
        self.slide.shapes.add_picture(str(f), IN((13.333 - w_in) / 2), IN(self.y), IN(w_in), IN(w_in * ratio))
        self.y += h

def title_slide(prs, title, subtitle, src):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, H - Inches(0.14), W, Inches(0.14), BLUE)
    tb = s.shapes.add_textbox(Inches(1.05), Inches(2.45), Inches(11.2), Inches(2.4))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title
    r.font.size = Pt(38); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    p2 = tf.add_paragraph(); p2.space_before = Pt(12)
    r2 = p2.add_run(); r2.text = subtitle
    r2.font.size = Pt(17); r2.font.color.rgb = BLUE; r2.font.name = FONT
    if src:
        p3 = tf.add_paragraph(); p3.space_before = Pt(24)
        r3 = p3.add_run(); r3.text = src
        r3.font.size = Pt(10.5); r3.font.color.rgb = RGBColor(0x8C, 0x9A, 0xAF); r3.font.name = FONT

def build(prs, md, root, course, kicker_prefix=""):
    title, src, sections = parse(md)
    title_slide(prs, title, course, src)
    deck = Deck(prs, f"{kicker_prefix}{title}")
    for sec in sections:
        deck.start_section(sec["title"] or title)
        for unit in group_units(sec["blocks"]):
            hs = [height_of(b, root) for b in unit]
            if sum(hs) <= BODY_H_IN:
                deck.ensure(sum(hs))
            for b, h in zip(unit, hs):
                kind, val = b
                if kind == "para": deck.para(val, h)
                elif kind == "list": deck.bullets(val, h)
                elif kind == "img": deck.image(val, h, root)
                elif kind == "code":
                    lines = val.split("\n")
                    if len(lines) <= CODE_MAX_LINES:
                        deck.code(val, h)
                    else:
                        for k in range(0, len(lines), CODE_MAX_LINES):
                            piece = "\n".join(lines[k:k + CODE_MAX_LINES])
                            deck.slide = None
                            deck.code(piece, height_of(["code", piece], root))

def main():
    a = sys.argv[1:]
    merge = "--merge" in a; a = [x for x in a if x != "--merge"]
    only = None
    if "--only" in a:
        k = a.index("--only"); only = set(a[k+1:]); a = a[:k]
    srcdir, outdir, course = a[0], a[1], a[2]
    root = Path(__file__).parent
    files = sorted(f for f in glob.glob(str(root / srcdir / "*.md")) if not f.endswith("README.md"))
    if only: files = [f for f in files if Path(f).name in only]
    out = root / outdir; out.mkdir(exist_ok=True)
    if merge:
        prs = Presentation(); prs.slide_width, prs.slide_height = W, H
        title_slide(prs, course, "28Tech · Trọn bộ bài giảng", "https://blog.28tech.com.vn")
        for i, f in enumerate(files, 1):
            build(prs, f, root / srcdir, course, kicker_prefix=f"Bài {i} · ")
            print(f"  + {Path(f).name}")
        dest = out / f"{outdir}-full.pptx"; prs.save(dest)
        print(f"\n{dest.name}: {len(prs.slides._sldIdLst)} slide")
    else:
        tot = 0
        for f in files:
            prs = Presentation(); prs.slide_width, prs.slide_height = W, H
            build(prs, f, root / srcdir, course)
            n = len(prs.slides._sldIdLst); tot += n
            dest = out / (Path(f).stem + ".pptx"); prs.save(dest)
            print(f"  ok {dest.name}  ({n} slide)")
        print(f"\n{len(files)} file pptx, {tot} slide -> {out}")

if __name__ == "__main__":
    main()
